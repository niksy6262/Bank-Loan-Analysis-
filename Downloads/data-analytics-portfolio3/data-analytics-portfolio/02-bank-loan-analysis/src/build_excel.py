"""
Build excel/loan_analysis.xlsx - an Excel version of the loan analysis.

The 'Data' sheet holds the cleaned loans; every number on the other sheets is a LIVE
formula (COUNTIFS / SUMIFS / AVERAGEIFS), so if you change the data the summary updates.
Sheets: Summary (KPIs) | By Credit Band | By DTI Band | By Purpose (each with a chart)

Run:     python src/build_excel.py        (run src/analysis.py first)
"""
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
CLEAN = ROOT / "data" / "processed" / "loan_data_clean.csv"
OUT = ROOT / "excel" / "loan_analysis.xlsx"
OUT.parent.mkdir(exist_ok=True)

df = pd.read_csv(CLEAN)
n = len(df)
LAST = n + 1                                  # last data row (row 1 = header)

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="2F5D8C")
thin = Side(style="thin", color="BFC5CC")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)


def rng(col):                                 # absolute range on the Data sheet, e.g. Data!$N$2:$N$10001
    return f"Data!${col}$2:${col}${LAST}"


# Column letters on the Data sheet (same order as the CSV)
C = {name: get_column_letter(i + 1) for i, name in enumerate(df.columns)}

wb = Workbook()

# ---------------------------------------------------------------- Data sheet
ws_data = wb.active
ws_data.title = "Data"
ws_data.append(list(df.columns))
for row in df.itertuples(index=False):
    ws_data.append(list(row))
for cell in ws_data[1]:
    cell.font = Font(name=FONT, bold=True, color="FFFFFF")
    cell.fill = HEAD_FILL
ws_data.freeze_panes = "A2"
ws_data.auto_filter.ref = f"A1:{get_column_letter(len(df.columns))}{LAST}"
for i, name in enumerate(df.columns, start=1):
    ws_data.column_dimensions[get_column_letter(i)].width = max(12, len(name) + 3)


def style_table(ws, first_row, last_row, ncols):
    for r in range(first_row, last_row + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(r, c)
            cell.border = BORDER
            cell.font = Font(name=FONT, bold=(r == first_row), color="FFFFFF" if r == first_row else "000000")
            if r == first_row:
                cell.fill = HEAD_FILL
                cell.alignment = Alignment(horizontal="center", wrap_text=True)


# ------------------------------------------------------------- Summary sheet
ws = wb.create_sheet("Summary", 0)
ws["A1"] = "Bank Loan Analysis - KPI Summary"
ws["A1"].font = Font(name=FONT, size=16, bold=True, color="2F5D8C")
ws["A2"] = "All values are formulas on the 'Data' sheet (synthetic data - see README)."
ws["A2"].font = Font(name=FONT, italic=True, color="666666")

ws.append([])
ws.append(["KPI", "Value"])
kpis = [
    ("Total applications", f"=COUNTA({rng(C['loan_id'])})", "#,##0"),
    ("Approved loans", f'=COUNTIF({rng(C["application_status"])},"Approved")', "#,##0"),
    ("Approval rate", "=B6/B5", "0.0%"),
    ("Total disbursed (INR)", f"=SUM({rng(C['disbursed_amount'])})", "#,##0"),
    ("Average interest rate (approved, %)", f'=AVERAGEIF({rng(C["application_status"])},"Approved",{rng(C["interest_rate"])})', "0.00"),
    ("Average DTI (approved, %)", f'=AVERAGEIF({rng(C["application_status"])},"Approved",{rng(C["dti_pct"])})', "0.0"),
    ("Good loans (Fully Paid + Current)", f'=COUNTIF({rng(C["loan_group"])},"Good Loan")', "#,##0"),
    ("Bad loans (Default)", f'=COUNTIF({rng(C["loan_group"])},"Bad Loan")', "#,##0"),
    ("Good loan % of approved", "=B11/B6", "0.0%"),
    ("Default rate (bad / approved)", "=B12/B6", "0.0%"),
    ("Amount at risk (defaulted amount / disbursed)",
     f'=SUMIF({rng(C["loan_group"])},"Bad Loan",{rng(C["disbursed_amount"])})/B8', "0.0%"),
]
for i, (label, formula, fmt) in enumerate(kpis, start=5):
    ws.cell(i, 1, label)
    ws.cell(i, 2, formula).number_format = fmt
style_table(ws, 4, 4 + len(kpis), 2)
for r in range(5, 5 + len(kpis)):
    ws.cell(r, 2).alignment = Alignment(horizontal="right")
ws.column_dimensions["A"].width = 48
ws.column_dimensions["B"].width = 20


# ------------------------------------------------- band / purpose sheets
def breakdown_sheet(title, group_col, groups, heading):
    sh = wb.create_sheet(title)
    sh["A1"] = title
    sh["A1"].font = Font(name=FONT, size=14, bold=True, color="2F5D8C")
    hdr = [heading, "Applications", "Approved", "Approval rate", "Bad loans", "Default rate", "Disbursed (INR)"]
    for j, h in enumerate(hdr, start=1):
        sh.cell(3, j, h)
    g = rng(C[group_col])
    for i, grp in enumerate(groups, start=4):
        sh.cell(i, 1, grp)
        sh.cell(i, 2, f"=COUNTIF({g},A{i})")
        sh.cell(i, 3, f'=COUNTIFS({g},A{i},{rng(C["application_status"])},"Approved")')
        sh.cell(i, 4, f"=IFERROR(C{i}/B{i},0)").number_format = "0.0%"
        sh.cell(i, 5, f'=COUNTIFS({g},A{i},{rng(C["loan_group"])},"Bad Loan")')
        sh.cell(i, 6, f"=IFERROR(E{i}/C{i},0)").number_format = "0.0%"
        sh.cell(i, 7, f"=SUMIF({g},A{i},{rng(C['disbursed_amount'])})").number_format = "#,##0"
    last = 3 + len(groups)
    style_table(sh, 3, last, len(hdr))
    for j in range(1, len(hdr) + 1):
        sh.column_dimensions[get_column_letter(j)].width = 18
    ch = BarChart()
    ch.type = "col"
    ch.title = f"Default rate by {heading.lower()}"
    ch.y_axis.title = "Default rate"
    ch.y_axis.numFmt = "0%"
    ch.add_data(Reference(sh, min_col=6, min_row=3, max_row=last), titles_from_data=True)
    ch.set_categories(Reference(sh, min_col=1, min_row=4, max_row=last))
    ch.legend = None
    ch.height, ch.width = 8, 16
    sh.add_chart(ch, f"A{last + 3}")


breakdown_sheet("By Credit Band", "credit_band",
                ["Below 600", "600-649", "650-699", "700-749", "750-799", "800+"], "Credit band")
breakdown_sheet("By DTI Band", "dti_band", ["Below 30%", "30-40%", "40-50%", "50%+"], "DTI band")
breakdown_sheet("By Purpose", "loan_purpose", ["Home", "Business", "Personal", "Auto", "Education"], "Loan purpose")

wb.save(OUT)
print("Saved", OUT)
